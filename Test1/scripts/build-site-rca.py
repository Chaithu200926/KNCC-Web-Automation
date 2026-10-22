"""
Builds the RCA workbook of one website run (WEB-01..48) from test-results-<site>/results.json, in the layout of the
earlier "Cinescape UAT vs Pre-UAT RCA" workbooks:
  sheet 1 "RCA Report"      summary only (counts are formulas over sheet 2),
  sheet 2 "All test cases"  every case: description, steps, result, issue type, reason, remarks, link to the dashboard,
  sheet 3 "Failed cases"    the failed cases with an empty "Manual check status" (Passed / Failed drop-down).
A second workbook "Cinescape <Site> vs UAT and Pre-UAT RCA (Run N)" compares this run with the newest reviewed
"UAT vs Pre-UAT RCA" workbook (--baseline to pick another): results side by side, New failure / Fixed / Still failing,
and a "Differences" sheet with a manual check column.

Issue type and reason are a first guess from the failure (please review them; the drop-down lets you change them):
  Application issue  a check of the website's behaviour failed (the test's own message says what was expected),
  Script issue       the test could not find or use a page element (locator / timeout), e.g. after a page change,
  Environment issue  the website or the payment gateway could not be reached,
  To be checked      anything else.

With --final it builds "Cinescape <Site> Final RCA - <date>.xlsx" (and the Final comparison) from the day's
"RCA (Run N)" workbooks instead: each case with its result from the latest run that ran it, hand corrections kept.

  python scripts/build-site-rca.py pre-prod [--final] [--out-dir "<folder>"] [--cases "<test cases workbook>"]
    --out-dir  default: C:\\softwares\\KNCC testing\\<d Mon> reports (the run number goes up per site and day)
    --cases    workbook with the Web sheet (Test ID, Module, Test case, Test steps, Expected result); default: the
               newest "*All Test Cases*.xlsx" under C:\\softwares\\KNCC testing. Without it the steps come from the run.
"""
import argparse
import datetime as dt
import glob
import json
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.datavalidation import DataValidation

SITES = {
    'uat': ('UAT', 'http://10.35.1.66:8103'),
    'pre-uat': ('Pre-UAT', 'http://10.35.1.66:8123'),
    'pre-prod': ('Pre-Prod', 'http://10.35.1.66:8113'),
}
ISSUE_TYPES = ['Application issue', 'Script issue', 'Environment issue', 'To be checked']
TESTING = r'C:\softwares\KNCC testing'
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # The Test1 folder.

parser = argparse.ArgumentParser()
parser.add_argument('site', choices=SITES)
parser.add_argument('--out-dir')
parser.add_argument('--cases')
parser.add_argument('--baseline')
parser.add_argument('--final', action='store_true',
                    help="merge today's 'RCA (Run N)' workbooks of the site: every case with its latest result")
args = parser.parse_args()
label, url = SITES[args.site]

results = {}
if not args.final:
    results_file = os.path.join(ROOT, f'test-results-{args.site}', 'results.json')
    if not os.path.exists(results_file):
        sys.exit(f'No test-results-{args.site}/results.json - run the tests first.')
    with open(results_file, encoding='utf8') as f:
        results = json.load(f)


# ---------- the run's results, one entry per WEB case ----------
def clean(text):
    return ILLEGAL_CHARACTERS_RE.sub('', re.sub(r'\x1b\[[0-9;]*m', '', str(text or ''))).strip()  # No colours or control characters.


def first_lines(text, count=3):
    return ' / '.join([l.strip() for l in clean(text).split('\n') if l.strip()][:count])


def failed_step(steps):
    """Title of the deepest test step with an error ("" when the error was outside the steps)."""
    for step in steps or []:
        if step.get('error'):
            return failed_step(step.get('steps')) or step['title']
    return ''


def guess(message):
    """(issue type, reason) for a failure message."""
    text = clean(message)
    first = re.sub(r'^Error:\s*', '', text.split('\n')[0]).strip()
    if re.search(r'net::ERR_|ECONNREFUSED|ECONNRESET|ENOTFOUND|Bad Gateway|Service Unavailable|HTTP 50\d', text):
        return 'Environment issue', 'Website or gateway not reachable'
    if re.match(r'Test timeout of', first):
        waiting = re.search(r'waiting for (.+)', text)
        return 'To be checked', 'Test timed out' + (f' waiting for {waiting.group(1)[:80]}' if waiting else '')
    # The tests' own messages ("The new first name should be saved ...") describe the website's behaviour.
    if first and not re.match(r'(expect\(|locator\.|page\.|Timeout|TimeoutError|Target page|strict mode|elementHandle|frame\.)', first):
        return 'Application issue', re.split(r'\s*\(|\.\s', first)[0][:120]
    element = re.search(r"(?:waiting for|Locator:)\s*(.+)", text)
    if element or 'strict mode violation' in text or re.match(r'(TimeoutError|locator\.|page\.)', first):
        return 'Script issue', 'Locator change - ' + (element.group(1).strip()[:90] if element else first[:90])
    return 'To be checked', first[:120]


cases = {}
dashboard_index = 0  # The dashboard numbers the tests in report order (#test-1, #test-2, ...).


def walk(suite):
    global dashboard_index
    for spec in suite.get('specs', []):
        for test in spec.get('tests', []):
            dashboard_index += 1
            match = re.match(r'(WEB-\d+)\s+(.*)', spec['title'])
            if not match:
                continue
            last = (test.get('results') or [{}])[-1]
            status = last.get('status', 'skipped')
            entry = {
                'id': match.group(1), 'title': match.group(2), 'anchor': dashboard_index,
                'result': {'passed': 'Passed', 'skipped': 'Skipped'}.get(status, 'Failed'),
                'steps': [s['title'] for s in last.get('steps', []) if s.get('category', 'test.step') == 'test.step'],
                'duration': round((last.get('duration') or 0) / 1000),
                'issue': '', 'reason': '', 'step': '', 'remarks': '',
            }
            if entry['result'] == 'Failed':
                message = (last.get('error') or {}).get('message') or first_lines(json.dumps(last.get('errors')))
                entry['issue'], entry['reason'] = guess(message)
                entry['step'] = failed_step(last.get('steps'))
                extra = len(last.get('errors') or []) - 1
                entry['remarks'] = first_lines(message, 4) + (f' (+{extra} more check(s) failed)' if extra > 0 else '')
            elif entry['result'] == 'Skipped':
                notes = [a.get('description', '') for a in test.get('annotations', []) if a.get('type') == 'skip']
                entry['remarks'] = '; '.join(n for n in notes if n) or 'Skipped'
            cases[entry['id']] = entry
    for child in suite.get('suites', []):
        walk(child)


for suite in results.get('suites', []):
    walk(suite)
if not cases and not args.final:
    sys.exit('No WEB cases in the results.')

# ---------- descriptions and steps from the test cases workbook (optional) ----------
cases_book = args.cases
if not cases_book:
    found = [p for p in glob.glob(os.path.join(TESTING, '**', '*All Test Cases*.xlsx'), recursive=True) if '~$' not in p]
    cases_book = max(found, key=os.path.getmtime) if found else None
sheet_info = {}
if cases_book and os.path.exists(cases_book):
    try:
        book = load_workbook(cases_book, read_only=True, data_only=True)
        web = book['Web']
        header = None
        for row in web.iter_rows(values_only=True):
            if header is None:
                if row and row[0] == 'Test ID':
                    header = [str(c or '').strip() for c in row]
                continue
            if row and isinstance(row[0], str) and row[0].startswith('WEB-'):
                sheet_info[row[0]] = dict(zip(header, row))
    except Exception as error:  # Excel has it open, or a different layout: use the run's own steps.
        print(f'Test cases workbook not used ({error}).')

# ---------- output file: <out dir>\Cinescape <Site> RCA (Run N) - <date>.xlsx ----------
started = results.get('stats', {}).get('startTime', '')
run_day = dt.datetime.fromisoformat(started.replace('Z', '+00:00')).astimezone() if started else dt.datetime.now()
date_text = f'{run_day.day} {run_day:%b %Y}'
out_dir = args.out_dir or os.path.join(TESTING, f'{run_day.day} {run_day:%b} reports')
os.makedirs(out_dir, exist_ok=True)
pattern = re.compile(rf'Cinescape {re.escape(label)} RCA \(Run (\d+)\) - {re.escape(date_text)}\.xlsx$')
run_files = {int(m.group(1)): os.path.join(out_dir, m.group(0)) for m in (pattern.match(n) for n in os.listdir(out_dir)) if m}
dashboard = os.path.join(ROOT, f'dashboard-{args.site}', 'index.html')
if not args.final:
    run_no = 1 + max(run_files or [0])
    run_name = f'Run {run_no}'
    out_file = os.path.join(out_dir, f'Cinescape {label} RCA ({run_name}) - {date_text}.xlsx')
else:
    # Final RCA: every case with its result from the latest of today's runs that ran it (a rerun of a few cases
    # replaces only those). Issue types and reasons are taken as they are in those workbooks, with any hand corrections.
    if not run_files:
        sys.exit(f'No "Cinescape {label} RCA (Run N) - {date_text}.xlsx" in {out_dir}.')
    latest_run = max(run_files)
    for number in sorted(run_files):
        rows = list(load_workbook(run_files[number], read_only=True, data_only=True)['All test cases'].iter_rows(values_only=True))
        head = next(r for r in rows if r and r[0] == 'Test ID')
        for r in rows:
            if not (r and isinstance(r[0], str) and r[0].startswith('WEB-')):
                continue
            v = dict(zip(head, r))
            cases[r[0]] = {
                'id': r[0], 'title': v.get('Test case') or '', 'run': number,
                'anchor': None,  # The dashboard shows only the latest run, so no link for older results.
                'result': v.get(f'{label} result') or 'Not run', 'steps': [], 'duration': v.get('Duration (s)'),
                'issue': v.get(f'{label} issue type') or '', 'reason': v.get(f'{label} reason') or '',
                'step': v.get('Failed at step') or '', 'remarks': f'[Run {number}] ' + (v.get('Remarks') or ''),
            }
    # Links into the dashboard for the cases of the latest run (same order as that run's report).
    order = [i for i in sorted(cases, key=lambda i: int(i.split('-')[1])) if cases[i]['run'] == latest_run]
    for anchor, case_id in enumerate(order, start=1):
        cases[case_id]['anchor'] = anchor
    run_no = latest_run
    run_name = f'Final (Runs {", ".join(map(str, sorted(run_files)))})'
    out_file = os.path.join(out_dir, f'Cinescape {label} Final RCA - {date_text}.xlsx')

# ---------- workbook ----------
BOLD, TITLE, GREY = Font(bold=True), Font(bold=True, size=13), Font(color='FF595959')
HEAD_FILL = PatternFill('solid', fgColor='FFF2F2F2')
WRAP = Alignment(wrap_text=True, vertical='top')
fills = {'Passed': 'FFC6EFCE', 'Failed': 'FFFFC7CE', 'Skipped': 'FFEDEDED', 'Application issue': 'FFFFC7CE',
         'Script issue': 'FFFFEB9C', 'Environment issue': 'FFDDEBF7', 'To be checked': 'FFEDEDED'}


def colour(ws, ref, values):
    for value in values:
        ws.conditional_formatting.add(ref, CellIsRule(operator='equal', formula=[f'"{value}"'], fill=PatternFill('solid', fgColor=fills[value])))


def header_row(ws, row, titles, widths):
    for col, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row, col, title)
        cell.font, cell.fill, cell.alignment = BOLD, HEAD_FILL, WRAP
        ws.column_dimensions[cell.column_letter].width = width


ids = sorted(cases, key=lambda i: int(i.split('-')[1]))
linked = lambda c: bool(c and c.get('anchor') and os.path.exists(dashboard))  # A link to the test in the dashboard.
wb = Workbook()

# Sheet 2 first (sheet 1's formulas point at it).
allc = wb.active
allc.title = 'All test cases'
cols = ['Test ID', 'Module', 'Test case', 'Description (expected result)', 'Test steps', f'{label} result',
        f'{label} issue type', f'{label} reason', 'Failed at step', 'Remarks', 'Duration (s)', 'Evidence']
header_row(allc, 3, cols, [9, 16, 30, 40, 55, 11, 18, 34, 34, 60, 10, 14])
allc.cell(1, 1, f'All test cases - {label} website, {run_name} ({date_text})').font = TITLE
allc.cell(2, 1, f'{label}: {url}    Issue type and reason are a first guess from the failure - please review.').font = GREY
for row, case_id in enumerate(ids, start=4):
    c, info = cases[case_id], sheet_info.get(case_id, {})
    steps = info.get('Test steps') or '\n'.join(f'{n}. {s}' for n, s in enumerate(c['steps'], start=1))
    values = [case_id, info.get('Module') or '', info.get('Test case') or c['title'], info.get('Expected result') or '',
              steps, c['result'], c['issue'] or None, c['reason'] or None, c['step'] or None, c['remarks'] or None,
              c['duration'], 'Dashboard' if linked(c) else None]
    for col, value in enumerate(values, start=1):
        allc.cell(row, col, value).alignment = WRAP
    if linked(c):
        allc.cell(row, 12).hyperlink = f'file:///{dashboard.replace(os.sep, "/")}#test-{c["anchor"]}'
        allc.cell(row, 12).font = Font(color='FF0563C1', underline='single')
last = 3 + len(ids)
allc.freeze_panes = 'C4'
allc.auto_filter.ref = f'A3:L{last}'
issue_list = DataValidation(type='list', formula1='"' + ','.join(ISSUE_TYPES) + '"', allow_blank=True)
allc.add_data_validation(issue_list)
issue_list.add(f'G4:G{last}')
colour(allc, f'F4:F{last}', ['Passed', 'Failed', 'Skipped'])
colour(allc, f'G4:G{last}', ISSUE_TYPES)

# Sheet 1: summary.
summary = wb.create_sheet('RCA Report', 0)
summary.column_dimensions['A'].width, summary.column_dimensions['B'].width = 34, 14
rng = lambda col: f"'All test cases'!${col}$4:${col}${last}"
summary.cell(1, 1, f'Cinescape website - RCA report ({label}) - {run_name}').font = TITLE
notes = [
    f'{label}: {url}    ' + (f'Latest result of each case from runs {", ".join(map(str, sorted(run_files)))} of {date_text} (Remarks start with the run).' if args.final else f'Run started {run_day:%d %b %Y %H:%M}, {len(ids)} cases, run on the QA PC.'),
    'Issue type: "Application issue" = the website needs a fix.  "Script issue" = the website changed and the test needs an update (e.g. Locator change).  "Environment issue" = site or gateway not reachable.',
    'Issue type and reason are filled automatically from the failure messages - please review them (drop-down on the sheet "All test cases").',
    'The full details of every test case are on the sheet "All test cases"; the failed cases with a manual check column are on "Failed cases".',
]
for row, note in enumerate(notes, start=2):
    summary.cell(row, 1, note).font = GREY
summary.cell(7, 2, label).font = BOLD
lines = [('Cases run', f'=COUNTA({rng("A")})-COUNTIF({rng("F")},"Skipped")'),
         ('Passed', f'=COUNTIF({rng("F")},"Passed")'),
         ('Failed', f'=COUNTIF({rng("F")},"Failed")'),
         ('Skipped', f'=COUNTIF({rng("F")},"Skipped")')]
lines += [(f'  Failed - {t}', f'=COUNTIFS({rng("F")},"Failed",{rng("G")},"{t}")') for t in ISSUE_TYPES]
lines += [('Pass rate (of cases run)', f'=IFERROR(B9/B8,0)')]
for row, (name, formula) in enumerate(lines, start=8):
    summary.cell(row, 1, name)
    summary.cell(row, 2, formula)
summary.cell(8 + len(lines) - 1, 2).number_format = '0%'

# Sheet 3: failed cases with a manual check column.
failed = wb.create_sheet('Failed cases')
fcols = ['Test ID', 'Test case', f'{label} result', f'{label} issue type', f'{label} reason', 'Failed at step', 'Remarks',
         'Evidence', 'Manual check status', 'Manual check notes']
header_row(failed, 3, fcols, [9, 30, 11, 18, 34, 34, 60, 14, 16, 40])
failed.cell(1, 1, f'Failed cases - {label} website, {run_name} ({date_text})').font = TITLE
failed.cell(2, 1, 'Manual check status: choose Passed or Failed after checking the case by hand.').font = GREY
row = 3
for case_id in ids:
    c = cases[case_id]
    if c['result'] != 'Failed':
        continue
    row += 1
    values = [case_id, sheet_info.get(case_id, {}).get('Test case') or c['title'], c['result'], c['issue'], c['reason'],
              c['step'] or None, c['remarks'], 'Dashboard' if linked(c) else None, None, None]
    for col, value in enumerate(values, start=1):
        failed.cell(row, col, value).alignment = WRAP
    if linked(c):
        failed.cell(row, 8).hyperlink = f'file:///{dashboard.replace(os.sep, "/")}#test-{c["anchor"]}'
        failed.cell(row, 8).font = Font(color='FF0563C1', underline='single')
if row == 3:
    failed.cell(4, 1, 'No failed cases in this run.')
failed.freeze_panes = 'C4'
end = max(row, 4)
manual = DataValidation(type='list', formula1='"Passed,Failed"', allow_blank=True)
issue_list2 = DataValidation(type='list', formula1='"' + ','.join(ISSUE_TYPES) + '"', allow_blank=True)
failed.add_data_validation(manual)
failed.add_data_validation(issue_list2)
manual.add(f'I4:I{end}')
issue_list2.add(f'D4:D{end}')
colour(failed, f'I4:I{end}', ['Passed', 'Failed'])
colour(failed, f'D4:D{end}', ISSUE_TYPES)

wb.save(out_file)
print(f'RCA workbook: {out_file}')

# ======================================================================================================================
# Comparison RCA: this run next to the newest reviewed "UAT vs Pre-UAT RCA" workbook (its sheet "All test cases" has
# "<site> result / issue type / reason" columns per site, checked by hand).
# ======================================================================================================================
baseline_file = args.baseline
if not baseline_file:
    found = [p for p in glob.glob(os.path.join(TESTING, '**', 'Cinescape UAT vs Pre-UAT RCA*.xlsx'), recursive=True) if '~$' not in p]
    baseline_file = max(found, key=os.path.getmtime) if found else None
if not baseline_file or not os.path.exists(baseline_file):
    print('No "UAT vs Pre-UAT RCA" workbook found - comparison skipped.')
    sys.exit(0)

base_sheet = load_workbook(baseline_file, read_only=True, data_only=True)['All test cases']
base_rows = list(base_sheet.iter_rows(values_only=True))
base_header = next(r for r in base_rows if r and r[0] == 'Test ID')
base_title = next((r[0] for r in base_rows if r and r[0]), '')  # e.g. "All test cases - UAT vs Pre-UAT, Run 2 (14 Oct 2026)"
base_when = (re.search(r'(Run \d+ \([^)]*\))', str(base_title)) or re.search(r'(.*)', os.path.basename(baseline_file))).group(1)
other_sites = [l for l in ('UAT', 'Pre-UAT', 'Pre-Prod') if l != label and f'{l} result' in base_header]
baseline = {}
for r in base_rows:
    if r and isinstance(r[0], str) and r[0].startswith('WEB-'):
        row_values = dict(zip(base_header, r))
        baseline[r[0]] = {l: (row_values.get(f'{l} result') or 'Not run', row_values.get(f'{l} issue type') or '',
                              row_values.get(f'{l} reason') or '') for l in other_sites}
compare_to = 'Pre-UAT' if 'Pre-UAT' in other_sites else (other_sites[0] if other_sites else None)


def change(mine, theirs):
    if {mine, theirs} & {None, 'Not run', 'Skipped'}:
        return 'Not compared'
    if mine == theirs:
        return 'Same - passed' if mine == 'Passed' else 'Still failing'
    return 'New failure' if mine == 'Failed' else 'Fixed'


all_ids = sorted(set(ids) | set(baseline), key=lambda i: int(i.split('-')[1]))
cmp_file = os.path.join(out_dir, f'Cinescape {label} vs {" and ".join(other_sites)} ' + ('Final RCA' if args.final else f'RCA ({run_name})') + f' - {date_text}.xlsx')
cwb = Workbook()
call = cwb.active
call.title = 'All test cases'
heads = ['Test ID', 'Test case']
widths = [9, 30]
for l in other_sites + [label]:
    heads += [f'{l} result', f'{l} issue type', f'{l} reason']
    widths += [11, 18, 30]
heads += [f'{label} vs {compare_to}', f'{label} failed at step', f'{label} remarks', 'Evidence']
widths += [16, 30, 60, 12]
header_row(call, 3, heads, widths)
call.cell(1, 1, f'All test cases - {" vs ".join(other_sites + [label])} ({label} {run_name}, {date_text})').font = TITLE
call.cell(2, 1, f'{" and ".join(other_sites)}: from "{os.path.basename(baseline_file)}" ({base_when}, reviewed).  '
                f'{label}: this run, issue type and reason guessed from the failure - please review.').font = GREY
change_col = 3 + 3 * (len(other_sites) + 1)
for row, case_id in enumerate(all_ids, start=4):
    c = cases.get(case_id)
    values = [case_id, sheet_info.get(case_id, {}).get('Test case') or (c['title'] if c else '')]
    for l in other_sites:
        values += list(baseline.get(case_id, {}).get(l, ('Not run', '', '')))
    mine = c['result'] if c else 'Not run'
    values += [mine, c['issue'] if c else '', c['reason'] if c else '']
    values += [change(mine, baseline.get(case_id, {}).get(compare_to, ('Not run',))[0]) if compare_to else 'Not compared',
               c['step'] if c else '', c['remarks'] if c else '', 'Dashboard' if linked(c) else None]
    for col, value in enumerate(values, start=1):
        call.cell(row, col, value or None).alignment = WRAP
    if linked(c):
        call.cell(row, len(values)).hyperlink = f'file:///{dashboard.replace(os.sep, "/")}#test-{c["anchor"]}'
        call.cell(row, len(values)).font = Font(color='FF0563C1', underline='single')
clast = 3 + len(all_ids)
call.freeze_panes = 'C4'
call.auto_filter.ref = f'A3:{call.cell(3, len(heads)).column_letter}{clast}'
fills.update({'New failure': 'FFFFC7CE', 'Fixed': 'FFC6EFCE', 'Still failing': 'FFFFEB9C', 'Same - passed': 'FFE2EFDA',
              'Not compared': 'FFEDEDED', 'Not run': 'FFEDEDED'})
letter = lambda col: call.cell(3, col).column_letter
for n in range(len(other_sites) + 1):
    colour(call, f'{letter(3 + 3 * n)}4:{letter(3 + 3 * n)}{clast}', ['Passed', 'Failed', 'Skipped', 'Not run'])
    colour(call, f'{letter(4 + 3 * n)}4:{letter(4 + 3 * n)}{clast}', ISSUE_TYPES)
colour(call, f'{letter(change_col)}4:{letter(change_col)}{clast}', ['New failure', 'Fixed', 'Still failing', 'Same - passed', 'Not compared'])
mine_issue = DataValidation(type='list', formula1='"' + ','.join(ISSUE_TYPES) + '"', allow_blank=True)
call.add_data_validation(mine_issue)
mine_issue.add(f'{letter(change_col - 2)}4:{letter(change_col - 2)}{clast}')

# Summary sheet: counts per site and the changes against the compared site.
csum = cwb.create_sheet('RCA Report', 0)
csum.column_dimensions['A'].width = 34
csum.cell(1, 1, f'Cinescape website - RCA comparison ({" vs ".join(other_sites + [label])}) - {label} {run_name}').font = TITLE
cnotes = [
    '    '.join(f'{l}: {SITES[k][1]}' for k, (l, _) in SITES.items() if l in other_sites + [label]),
    f'{" and ".join(other_sites)} results, issue types and reasons: {os.path.basename(baseline_file)} ({base_when}).',
    (f'{label}: latest result of each case from the runs of {date_text}' if args.final else f'{label}: this run ({run_day:%d %b %Y %H:%M})') + f'; issue type and reason are a first guess from the failure - please review.',
    f'"{label} vs {compare_to}": New failure = fails on {label} but passed on {compare_to}; Fixed = the other way round.',
]
for row, note in enumerate(cnotes, start=2):
    csum.cell(row, 1, note).font = GREY
crng = lambda col: f"'All test cases'!${letter(col)}$4:${letter(col)}${clast}"
for n, l in enumerate(other_sites + [label]):
    col = 2 + n
    csum.column_dimensions[csum.cell(7, col).column_letter].width = 14
    csum.cell(7, col, l).font = BOLD
    res, iss = 3 + 3 * n, 4 + 3 * n
    rows = [f'=COUNTA({crng(1)})-COUNTIF({crng(res)},"Skipped")-COUNTIF({crng(res)},"Not run")',
            f'=COUNTIF({crng(res)},"Passed")', f'=COUNTIF({crng(res)},"Failed")', f'=COUNTIF({crng(res)},"Skipped")']
    rows += [f'=COUNTIFS({crng(res)},"Failed",{crng(iss)},"{t}")' for t in ISSUE_TYPES]
    for r, formula in enumerate(rows, start=8):
        csum.cell(r, col, formula)
for r, name in enumerate(['Cases run', 'Passed', 'Failed', 'Skipped'] + [f'  Failed - {t}' for t in ISSUE_TYPES], start=8):
    csum.cell(r, 1, name)
csum.cell(17, 1, f'{label} vs {compare_to}').font = BOLD
for r, name in enumerate(['New failure', 'Fixed', 'Still failing', 'Same - passed', 'Not compared'], start=18):
    csum.cell(r, 1, f'  {name}')
    csum.cell(r, 2, f'=COUNTIF({crng(change_col)},"{name}")')

# Differences sheet: the cases whose result is not the same everywhere, with a manual check column.
cdiff = cwb.create_sheet('Differences')
dheads = ['Test ID', 'Test case'] + [f'{l} result' for l in other_sites] + [
    f'{label} result', f'{label} vs {compare_to}', f'{label} issue type', f'{label} reason', f'{label} remarks',
    f'{" / ".join(other_sites)} reason', 'Manual check status', 'Manual check notes']
header_row(cdiff, 3, dheads, [9, 30] + [11] * len(other_sites) + [11, 16, 18, 30, 55, 40, 16, 40])
cdiff.cell(1, 1, f'Cases with different results - {" vs ".join(other_sites + [label])} ({date_text})').font = TITLE
cdiff.cell(2, 1, 'Every case whose result is not the same on all websites. Manual check status: choose Passed or Failed after checking by hand.').font = GREY
row = 3
for case_id in all_ids:
    c = cases.get(case_id)
    mine = c['result'] if c else 'Not run'
    others = [baseline.get(case_id, {}).get(l, ('Not run', '', '')) for l in other_sites]
    if all(o[0] == mine for o in others):
        continue
    row += 1
    values = [case_id, sheet_info.get(case_id, {}).get('Test case') or (c['title'] if c else '')] + [o[0] for o in others] + [
        mine, change(mine, baseline.get(case_id, {}).get(compare_to, ('Not run',))[0]) if compare_to else 'Not compared',
        c['issue'] if c else '', c['reason'] if c else '', c['remarks'] if c else '',
        ' / '.join(f'{l}: {o[2]}' for l, o in zip(other_sites, others) if o[2]), None, None]
    for col, value in enumerate(values, start=1):
        cdiff.cell(row, col, value or None).alignment = WRAP
if row == 3:
    cdiff.cell(4, 1, 'All cases have the same result on every website.')
dend = max(row, 4)
dl = lambda col: cdiff.cell(3, col).column_letter
n_other = len(other_sites)
for col in range(3, 4 + n_other):
    colour(cdiff, f'{dl(col)}4:{dl(col)}{dend}', ['Passed', 'Failed', 'Skipped', 'Not run'])
colour(cdiff, f'{dl(4 + n_other)}4:{dl(4 + n_other)}{dend}', ['New failure', 'Fixed', 'Still failing', 'Not compared'])
colour(cdiff, f'{dl(5 + n_other)}4:{dl(5 + n_other)}{dend}', ISSUE_TYPES)
check = DataValidation(type='list', formula1='"Passed,Failed"', allow_blank=True)
cdiff.add_data_validation(check)
check.add(f'{dl(len(dheads) - 1)}4:{dl(len(dheads) - 1)}{dend}')
colour(cdiff, f'{dl(len(dheads) - 1)}4:{dl(len(dheads) - 1)}{dend}', ['Passed', 'Failed'])
cdiff.freeze_panes = 'C4'

cwb.save(cmp_file)
print(f'Comparison RCA workbook: {cmp_file}')
